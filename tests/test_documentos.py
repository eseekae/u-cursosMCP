"""Tests de extracción de texto: archivos armados en memoria, sin U-Cursos."""

import io
import zipfile

import pytest

from ucursos_mcp import documentos


def _pdf_con_texto(texto: str) -> bytes:
    """PDF mínimo de una página con una línea de texto (Helvetica)."""
    stream = f"BT /F1 12 Tf 72 720 Td ({texto}) Tj ET".encode()
    objetos = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = io.BytesIO()
    out.write(b"%PDF-1.4\n")
    offsets = []
    for i, obj in enumerate(objetos, start=1):
        offsets.append(out.tell())
        out.write(b"%d 0 obj\n" % i + obj + b"\nendobj\n")
    xref = out.tell()
    out.write(b"xref\n0 %d\n0000000000 65535 f \n" % (len(objetos) + 1))
    for o in offsets:
        out.write(b"%010d 00000 n \n" % o)
    out.write(b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (len(objetos) + 1, xref))
    return out.getvalue()


def test_pdf_con_paginas_numeradas():
    tipo, texto = documentos.extraer_texto("clase1.pdf", _pdf_con_texto("Teorema de Green"))
    assert tipo == "pdf"
    assert "--- Página 1 ---" in texto
    assert "Teorema de Green" in texto


def test_pdf_escaneado_avisa():
    from pypdf import PdfWriter

    w = PdfWriter()
    w.add_blank_page(width=612, height=792)
    buf = io.BytesIO()
    w.write(buf)
    with pytest.raises(documentos.DocumentoNoLegible, match="escaneado"):
        documentos.extraer_texto("scan.pdf", buf.getvalue())


def test_docx_con_tabla():
    from docx import Document

    doc = Document()
    doc.add_paragraph("Pauta Control 1")
    tabla = doc.add_table(rows=1, cols=2)
    tabla.rows[0].cells[0].text = "P1"
    tabla.rows[0].cells[1].text = "2 pts"
    buf = io.BytesIO()
    doc.save(buf)
    tipo, texto = documentos.extraer_texto("pauta.docx", buf.getvalue())
    assert tipo == "docx"
    assert "Pauta Control 1" in texto
    assert "P1 | 2 pts" in texto


def test_pptx_con_notas():
    from pptx import Presentation

    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Derivadas"
    slide.notes_slide.notes_text_frame.text = "Ojo con la regla de la cadena"
    buf = io.BytesIO()
    prs.save(buf)
    tipo, texto = documentos.extraer_texto("clase.pptx", buf.getvalue())
    assert tipo == "pptx"
    assert "--- Diapositiva 1 ---" in texto
    assert "Derivadas" in texto
    assert "(Notas) Ojo con la regla de la cadena" in texto


def test_tipo_por_contenido_sin_extension():
    assert documentos.tipo(None, _pdf_con_texto("x")) == "pdf"


def test_texto_latin1():
    _, texto = documentos.extraer_texto("enunciado.txt", "Ecuación".encode("latin-1"))
    assert texto == "Ecuación"


def test_zip_lista_contenido():
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w") as z:
        z.writestr("tarea.py", "print(1)")
    with pytest.raises(documentos.DocumentoNoLegible, match="tarea.py"):
        documentos.extraer_texto("codigo.zip", buf.getvalue())


def test_office_antiguo_avisa():
    with pytest.raises(documentos.DocumentoNoLegible, match="antiguo"):
        documentos.extraer_texto("apunte.doc", b"\xd0\xcf\x11\xe0")
