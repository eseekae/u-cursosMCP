"""
Capa de documentos: archivo descargado -> texto que Claude pueda leer.

No sabe nada de U-Cursos: recibe bytes y un nombre de archivo. Soporta PDF, Word (.docx),
PowerPoint (.pptx) y archivos de texto o código. Los formatos viejos (.doc, .ppt) y los
PDF escaneados (solo imagen) no tienen texto extraíble y se avisa en vez de inventar.
"""

from __future__ import annotations

import io
import zipfile
from pathlib import PurePath


class DocumentoNoLegible(ValueError):
    """El archivo no tiene texto que se pueda extraer (formato no soportado, escaneado, etc.)."""


EXT_TEXTO = {
    ".txt", ".md", ".csv", ".tex", ".json", ".xml", ".html", ".sql",
    ".py", ".java", ".c", ".h", ".cpp", ".hpp", ".js", ".ts", ".m", ".r", ".jl", ".ipynb",
}


def tipo(nombre: str | None, contenido: bytes) -> str:
    """Detecta el tipo por extensión y, si no basta, por los primeros bytes."""
    ext = PurePath(nombre or "").suffix.lower()
    if ext in (".pdf", ".docx", ".pptx", ".zip"):
        return ext[1:]
    if ext in EXT_TEXTO:
        return "texto"
    if ext in (".doc", ".ppt", ".xls"):
        return "office_antiguo"
    if contenido.startswith(b"%PDF"):
        return "pdf"
    if contenido.startswith(b"PK"):
        try:
            nombres = zipfile.ZipFile(io.BytesIO(contenido)).namelist()
        except zipfile.BadZipFile:
            return "desconocido"
        if "word/document.xml" in nombres:
            return "docx"
        if "ppt/presentation.xml" in nombres:
            return "pptx"
        return "zip"
    return "desconocido"


def _pdf(contenido: bytes) -> str:
    from pypdf import PdfReader

    textos = [(p.extract_text() or "").strip() for p in PdfReader(io.BytesIO(contenido)).pages]
    if not any(textos):
        raise DocumentoNoLegible("El PDF no tiene texto: parece escaneado (solo imágenes).")
    return "\n\n".join(f"--- Página {i} ---\n{t}" for i, t in enumerate(textos, start=1))


def _docx(contenido: bytes) -> str:
    from docx import Document

    doc = Document(io.BytesIO(contenido))
    partes = [p.text for p in doc.paragraphs if p.text.strip()]
    for tabla in doc.tables:
        for fila in tabla.rows:
            partes.append(" | ".join(c.text.strip() for c in fila.cells))
    return "\n".join(partes)


def _pptx(contenido: bytes) -> str:
    from pptx import Presentation

    partes = []
    for i, slide in enumerate(Presentation(io.BytesIO(contenido)).slides, start=1):
        textos = [s.text_frame.text for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
        if slide.has_notes_slide and slide.notes_slide.notes_text_frame.text.strip():
            textos.append(f"(Notas) {slide.notes_slide.notes_text_frame.text}")
        partes.append(f"--- Diapositiva {i} ---\n" + "\n".join(textos))
    return "\n\n".join(partes)


def _texto(contenido: bytes) -> str:
    try:
        return contenido.decode("utf-8")
    except UnicodeDecodeError:
        return contenido.decode("latin-1")


def extraer_texto(nombre: str | None, contenido: bytes) -> tuple[str, str]:
    """Devuelve (tipo, texto). Lanza DocumentoNoLegible si no hay texto que sacar."""
    t = tipo(nombre, contenido)
    if t == "pdf":
        return t, _pdf(contenido)
    if t == "docx":
        return t, _docx(contenido)
    if t == "pptx":
        return t, _pptx(contenido)
    if t == "texto":
        return t, _texto(contenido)
    if t == "zip":
        nombres = zipfile.ZipFile(io.BytesIO(contenido)).namelist()
        raise DocumentoNoLegible("Es un ZIP; no se lee por dentro. Contiene: " + ", ".join(nombres[:50]))
    if t == "office_antiguo":
        raise DocumentoNoLegible("Formato de Office antiguo (.doc/.ppt/.xls): no se puede leer. Descárgalo con descargar_material.")
    raise DocumentoNoLegible("No reconozco el formato del archivo. Descárgalo con descargar_material.")
